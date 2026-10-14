"""
Pipeline de indexación incremental con CocoIndex.

    Fuente (RAG_SOURCE_DIR) -> extraer texto (PDF/código/docs) -> chunking
    -> embeddings (Ollama: nomic-embed-text) -> Qdrant

Uso (desde la raíz del repo):
    uv run cocoindex update rag.indexer      # sincroniza una vez (sólo procesa cambios)
    uv run cocoindex update -L rag.indexer   # modo live: vigila la carpeta y re-indexa al vuelo
"""

from __future__ import annotations

import asyncio
import io
import pathlib
import re
from typing import AsyncIterator

import cocoindex as coco
from cocoindex.connectors import localfs, qdrant
from cocoindex.ops.litellm import LiteLLMEmbedder
from cocoindex.ops.text import RecursiveSplitter, detect_code_language
from cocoindex.resources.chunk import Chunk
from cocoindex.resources.file import FileLike, PatternFilePathMatcher
from cocoindex.resources.id import IdGenerator
from qdrant_client import QdrantClient

from rag import config

QDRANT_DB = coco.ContextKey[QdrantClient]("rag_qdrant")
EMBEDDER = coco.ContextKey[LiteLLMEmbedder]("rag_embedder", detect_change=True)

_splitter = RecursiveSplitter()


@coco.lifespan
async def coco_lifespan(builder: coco.EnvironmentBuilder) -> AsyncIterator[None]:
    builder.provide(QDRANT_DB, qdrant.create_client(config.QDRANT_URL, prefer_grpc=True))
    builder.provide(
        EMBEDDER,
        LiteLLMEmbedder(
            f"ollama/{config.EMBED_MODEL}",
            api_base=config.OLLAMA_URL,
            max_inflight_requests=16,  # Concurrencia acelerada aprovechando la VRAM
        ),
    )
    yield


def _extract_pdf_pages(data: bytes) -> list[tuple[int, str]]:
    import pymupdf  # Motor en C/C++ ultrarrápido

    doc = pymupdf.open(stream=data, filetype="pdf")
    pages: list[tuple[int, str]] = []
    for i, page in enumerate(doc, start=1):
        try:
            text = page.get_text() or ""
        except Exception:
            continue
        if text.strip():
            pages.append((i, text))
    return pages


def _group_pdf_pages(pages: list[tuple[int, str]]) -> list[tuple[int, int, str]]:
    """Agrupa páginas consecutivas hasta ~CHUNK_SIZE bytes.

    En diapositivas cada página trae muy poco texto; indexarlas por separado produce
    fragmentos casi vacíos (sólo el título) que dominan la búsqueda léxica.
    Devuelve (página_inicial, página_final, texto).
    """
    groups: list[tuple[int, int, str]] = []
    buf: list[str] = []
    size = 0
    start = end = 0
    for page, text in pages:
        text = text.strip()
        text_size = len(text.encode())
        if buf and size + text_size > config.CHUNK_SIZE:
            groups.append((start, end, "\n\n".join(buf)))
            buf, size = [], 0
        if not buf:
            start = page
        buf.append(text)
        size += text_size
        end = page
    if buf:
        groups.append((start, end, "\n\n".join(buf)))
    return groups


def _extract_docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def _extract_notebook_text(raw: str) -> str:
    """Extrae sólo el código y el markdown de un .ipynb (sin metadatos ni salidas en base64)."""
    import json

    try:
        nb = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    cells = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else src
        if src.strip():
            cells.append(src if cell.get("cell_type") == "code" else f'"""\n{src}\n"""')
    return "\n\n".join(cells)


_HEADING_RE = {
    "typst": re.compile(r"^(=+)\s+(.+?)\s*$"),
    "markdown": re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$"),
}


def _headings(text: str, language: str | None) -> list[tuple[int, int, str]]:
    """Títulos del documento como (línea, nivel, título), con líneas desde 1."""
    pattern = _HEADING_RE.get(language or "")
    if pattern is None:
        return []
    found = []
    in_code = False
    for line_no, line in enumerate(text.splitlines(), start=1):
        if line.lstrip().startswith("```"):
            in_code = not in_code
            continue
        m = None if in_code else pattern.match(line)
        if m:
            found.append((line_no, len(m.group(1)), m.group(2)))
    return found


def _section_at(headings: list[tuple[int, int, str]], line: int) -> str | None:
    """Ruta de títulos vigente en una línea, p.ej. "Decisiones > ADR-04: Kafka"."""
    stack: list[tuple[int, str]] = []
    for line_no, level, title in headings:
        if line_no > line:
            break
        while stack and stack[-1][0] >= level:
            stack.pop()
        stack.append((level, title))
    return " > ".join(title for _, title in stack) or None


def _strip_trailing_headings(text: str, language: str | None) -> tuple[str, int]:
    """Quita títulos sueltos al final de un fragmento: pertenecen al fragmento siguiente
    y hacen que el LLM atribuya este contenido a la sección equivocada.
    Devuelve el texto y cuántas líneas se quitaron."""
    pattern = _HEADING_RE.get(language or "")
    if pattern is None:
        return text, 0
    lines = text.rstrip().splitlines()
    removed = 0
    while len(lines) > 1 and (not lines[-1].strip() or pattern.match(lines[-1])):
        lines.pop()
        removed += 1
    return "\n".join(lines), removed


@coco.fn
async def process_chunk(
    chunk: Chunk,
    rel_path: str,
    source: str,
    language: str | None,
    page: int | None,
    page_end: int | None,
    headings: list[tuple[int, int, str]],
    id_gen: IdGenerator,
    target: qdrant.CollectionTarget,
) -> None:
    text, removed = _strip_trailing_headings(chunk.text, language)
    section = _section_at(headings, chunk.start.line)
    # La ruta y la sección dan contexto a fragmentos que por sí solos son ambiguos
    # ("Parcial 1: 20%", "Consecuencias positivas: ...").
    embedding = await coco.use_context(EMBEDDER).embed(
        config.DOC_PREFIX + config.chunk_header(rel_path, section) + text
    )
    target.declare_point(
        qdrant.PointStruct(
            id=await id_gen.next_id(text),
            vector=embedding.tolist(),
            payload={
                "path": rel_path,
                "source": source,
                "language": language,
                "page": page,
                "page_end": page_end,
                "section": section,
                "start_line": chunk.start.line,
                "end_line": chunk.end.line - removed,
                "text": text,
            },
        )
    )


@coco.fn(memo=True)  # memo: si el archivo no cambió, no se vuelve a procesar
async def process_file(file: FileLike, target: qdrant.CollectionTarget) -> None:
    path = file.file_path.path
    # walk_dir entrega rutas absolutas cuando la fuente es un Path; se guardan relativas
    # a RAG_SOURCE_DIR para que "source" sea la carpeta de primer nivel y no "/".
    if path.is_absolute():
        for base in (config.SOURCE_DIR, config.SOURCE_DIR.resolve()):
            if path.is_relative_to(base):
                path = path.relative_to(base)
                break
    rel_path = path.as_posix()
    # "source" = primera carpeta dentro de RAG_SOURCE_DIR (p.ej. el nombre del repo o materia)
    source = path.parts[0] if len(path.parts) > 1 else "(raíz)"
    suffix = path.suffix.lower()
    is_pdf = suffix == ".pdf"
    is_docx = suffix == ".docx"

    size = await file.size()
    if size > (config.MAX_PDF_FILE_BYTES if (is_pdf or is_docx) else config.MAX_TEXT_FILE_BYTES):
        return

    # Secciones a trocear: (página inicial, página final, texto)
    sections: list[tuple[int | None, int | None, str]]
    if is_pdf:
        pages = await asyncio.to_thread(_extract_pdf_pages, await file.read())
        sections = list(_group_pdf_pages(pages))
        language = None
    elif is_docx:
        text = await asyncio.to_thread(_extract_docx_text, await file.read())
        sections = [(None, None, text)]
        language = None
    else:
        text = await file.read_text()
        if "\x00" in text:  # binario disfrazado
            return
        if suffix == ".ipynb":
            text = _extract_notebook_text(text)
            language = "python"
        else:
            language = detect_code_language(filename=path.name)
        sections = [(None, None, text)]

    split_language = language or ("markdown" if suffix in (".md", ".mdx") else None)
    id_gen = IdGenerator()
    for page, page_end, text in sections:
        chunks = _splitter.split(
            text,
            chunk_size=config.CHUNK_SIZE,
            min_chunk_size=config.MIN_CHUNK_SIZE,
            chunk_overlap=config.CHUNK_OVERLAP,
            language=split_language,
        )
        chunks = [c for c in chunks if c.text.strip()]
        headings = _headings(text, split_language)
        await coco.map(
            process_chunk, chunks, rel_path, source, split_language, page, page_end, headings, id_gen, target
        )


@coco.fn
async def app_main(source_dir: pathlib.Path) -> None:
    if not source_dir.is_dir():
        raise SystemExit(
            f"La carpeta de fuentes no existe: {source_dir}\n"
            "Créala (o define RAG_SOURCE_DIR) y pon ahí tus repos/documentos."
        )

    collection = await qdrant.mount_collection_target(
        QDRANT_DB,
        collection_name=config.QDRANT_COLLECTION,
        schema=await qdrant.CollectionSchema.create(
            vectors=qdrant.QdrantVectorDef(schema=EMBEDDER, distance="cosine")
        ),
    )
    files = localfs.walk_dir(
        source_dir,
        recursive=True,
        path_matcher=PatternFilePathMatcher(
            included_patterns=config.INCLUDED_PATTERNS,
            excluded_patterns=config.EXCLUDED_PATTERNS,
        ),
        live=True,  # sólo vigila si se ejecuta con `cocoindex update -L`
    )
    await coco.mount_each(process_file, files.items(), collection)


app = coco.App(
    coco.AppConfig(name="RagHomelab", max_inflight_components=8),
    app_main,
    source_dir=config.SOURCE_DIR,
)
